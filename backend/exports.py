from __future__ import annotations

import io
import json

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas


def _add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def report_to_docx(report: dict) -> bytes:
    document = Document()
    document.styles["Normal"].font.name = "宋体"
    document.styles["Normal"].font.size = Pt(10.5)
    document.add_heading("悟真鉴引 · 引文核验报告", 0)
    document.add_paragraph(f"来源：{report['source_name']}")
    document.add_paragraph(f"模式：{'文件模式' if report['mode'] == 'file' else '粘贴模式'}")
    document.add_heading("结论摘要", 1)
    summary = report["summary"]
    document.add_paragraph(
        f"文献总数 {summary['total']}；已证实 {summary['verified']}；"
        f"信息冲突 {summary['conflict']}；暂未检出 {summary['not_detected']}；"
        f"证据不足 {summary['insufficient_evidence']}。"
    )
    document.add_paragraph(report["narrative"]["overview"])
    document.add_heading("优先处理", 1)
    for issue in report["issues"]:
        document.add_paragraph(
            f"{issue['severity_label']}｜{issue['type']}｜{issue['message']} 建议：{issue['action']}",
            style="List Bullet",
        )
    document.add_heading("逐条文献", 1)
    for reference in report["references"]:
        _add_heading(document, f"{reference['ref_no']}. {reference['title'] or '未提供标题'}", 2)
        retraction_note = (
            (reference.get("retraction_note") or "未发现撤稿记录")
            if reference.get("retracted")
            else "未发现撤稿记录"
        )
        for label, value in (
            ("状态", reference["status_label"]),
            ("作者", reference["authors"] or "未提供"),
            ("年份", reference["year"] or "未提供"),
            ("DOI", reference["doi"] or "未提供"),
            ("撤稿状态", retraction_note),
            ("风险说明", reference["risk_reason"]),
            ("修复建议", reference["repair_suggestion"] or "暂无"),
        ):
            document.add_paragraph(f"{label}：{value}")
    document.add_heading("正文—证据链", 1)
    for citation in report["citations"]:
        document.add_paragraph(
            f"{citation['marker']}｜{citation['ref_title'] or '未匹配'}｜"
            f"{citation['support_label']}｜{citation['reason']}",
            style="List Bullet",
        )
        document.add_paragraph(f"正文：{citation['claim']}")
        if citation["evidence_quote"]:
            document.add_paragraph(f"证据：{citation['evidence_quote']}")
    document.add_heading("边界说明", 1)
    for boundary in report["boundaries"]:
        document.add_paragraph(boundary, style="List Bullet")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def report_to_xlsx(report: dict) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "核验结论"
    summary_rows = [
        ("项目", "数值"),
        ("文献总数", report["summary"]["total"]),
        ("已证实", report["summary"]["verified"]),
        ("信息冲突", report["summary"]["conflict"]),
        ("暂未检出", report["summary"]["not_detected"]),
        ("证据不足", report["summary"]["insufficient_evidence"]),
        ("正文引用", report["summary"]["citations"]),
        ("未匹配引用", report["summary"]["unresolved_citations"]),
    ]
    for row in summary_rows:
        summary_sheet.append(row)
    for cell in summary_sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EEF7")

    issues_sheet = workbook.create_sheet("问题清单")
    issues_sheet.append(["风险", "类型", "文献编号", "问题", "建议"])
    for issue in report["issues"]:
        issues_sheet.append(
            [
                issue["severity_label"],
                issue["type"],
                issue.get("ref_no", ""),
                issue["message"],
                issue["action"],
            ]
        )
    references_sheet = workbook.create_sheet("文献明细")
    references_sheet.append(
        ["编号", "标题", "作者", "年份", "DOI", "核验状态", "撤稿状态", "风险说明", "修复建议", "证据链接"]
    )
    for reference in report["references"]:
        references_sheet.append(
            [
                reference["ref_no"],
                reference["title"],
                reference["authors"],
                reference["year"],
                reference["doi"],
                reference["status_label"],
                reference.get("retraction_note", "") if reference.get("retracted") else "未发现",
                reference["risk_reason"],
                reference["repair_suggestion"],
                reference["evidence_url"],
            ]
        )
    citations_sheet = workbook.create_sheet("正文证据链")
    citations_sheet.append(["引用", "文献编号", "目标文献", "支持判断", "正文主张", "证据原文", "理由"])
    for citation in report["citations"]:
        citations_sheet.append(
            [
                citation["marker"],
                citation["ref_no"],
                citation["ref_title"],
                citation["support_label"],
                citation["claim"],
                citation["evidence_quote"],
                citation["reason"],
            ]
        )
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(48, max(12, max(len(str(cell.value or "")) for cell in column[:20]) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
            for cell in column:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def report_to_json(report: dict) -> bytes:
    public = {key: value for key, value in report.items() if key != "markdown"}
    return json.dumps(public, ensure_ascii=False, indent=2).encode("utf-8")


def report_to_pdf(report: dict) -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    output = io.BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    width, height = A4
    x, y = 42, height - 52

    def write(text: str, size: int = 10, step: int = 16) -> None:
        nonlocal y
        pdf.setFont("STSong-Light", size)
        for line in str(text).splitlines() or [""]:
            while line:
                chunk, line = line[:58], line[58:]
                if y < 50:
                    pdf.showPage()
                    pdf.setFont("STSong-Light", size)
                    y = height - 52
                pdf.drawString(x, y, chunk)
                y -= step

    write("悟真鉴引 · 引文核验报告", 18, 28)
    write(report["narrative"]["headline"], 12, 20)
    write(report["narrative"]["overview"])
    write("问题清单", 13, 22)
    for issue in report["issues"]:
        write(f"[{issue['severity_label']}] {issue['message']} 建议：{issue['action']}")
    write("逐条文献", 13, 22)
    for reference in report["references"]:
        retraction = " | 已撤稿" if reference.get("retracted") else ""
        write(
            f"{reference['ref_no']}. {reference['title'] or '未提供标题'} | "
            f"{reference['status_label']}{retraction} | {reference['risk_reason']}"
        )
    write("边界说明", 13, 22)
    for boundary in report["boundaries"]:
        write(f"• {boundary}")
    pdf.save()
    return output.getvalue()

