from __future__ import annotations

import io
import re
from dataclasses import dataclass

from docx import Document
from pypdf import PdfReader


class DocumentParseError(ValueError):
    pass


@dataclass
class ExtractedDocument:
    text: str
    file_type: str
    page_count: int = 0
    warning: str = ""


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("无法识别文本编码，请使用 UTF-8 或 GB18030 重新保存文件。")


def extract_document(filename: str, data: bytes) -> ExtractedDocument:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError("仅支持 PDF、DOCX、TXT 和 Markdown 文件。")

    if suffix == ".pdf":
        try:
            reader = PdfReader(io.BytesIO(data))
        except Exception as exc:
            raise DocumentParseError(f"PDF 无法打开：{exc}") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
        if len(text.strip()) < 30:
            raise DocumentParseError(
                "该 PDF 没有可提取的文字层，可能是扫描版。请上传带文本层的 PDF，或先使用 OCR 转换。"
            )
        return ExtractedDocument(text=text, file_type="pdf", page_count=len(reader.pages))

    if suffix == ".docx":
        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:
            raise DocumentParseError(f"DOCX 无法打开：{exc}") from exc
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                blocks.append("\t".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(blocks)
        if len(text.strip()) < 20:
            raise DocumentParseError("DOCX 中没有可提取文字。")
        return ExtractedDocument(text=text, file_type="docx")

    text = _decode_text(data)
    if len(text.strip()) < 20:
        raise DocumentParseError("文件内容为空或过短。")
    return ExtractedDocument(text=text, file_type=suffix.lstrip("."))


HEADING_RE = re.compile(
    r"(?mi)^[ \t]*(?:#{1,6}[ \t]*)?"
    r"(?:(?:[一二三四五六七八九十]+[、.．]|\d+[.、]?)[ \t]*)?"
    r"(?:参考文献|參考文獻|引用文献|引用文獻|主要参考文献|主要參考文獻|参考书目|參考書目|"
    r"references?(?:[ \t]+and[ \t]+notes)?|reference[ \t]+list|bibliography|"
    r"works[ \t]+cited|literature[ \t]+cited)"
    r"(?:[：:])?[ \t]*$"
)


def split_reference_section(text: str) -> tuple[str, str]:
    match = HEADING_RE.search(text)
    if not match:
        raise DocumentParseError(
            "未检测到有效的“参考文献 / References”分节，请补充标题，或直接复制文献列表进行检测。"
        )
    body = text[: match.start()].strip()
    references = text[match.end() :].strip()
    if not references:
        raise DocumentParseError(
            "已找到“参考文献”标题，但标题后没有可识别内容。请检查条目格式或粘贴文献列表。"
        )
    return body, references

